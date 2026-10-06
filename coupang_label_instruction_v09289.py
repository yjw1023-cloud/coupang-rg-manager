"""RG Manager v0.9.295 — Coupang label work-instruction workbook generator.

User uploads the edited China purchasing/order workbook. The page extracts the
rows, links them to ERP RG/BOM products, lets the user correct option IDs / label
quantities / packing instructions, and downloads a warehouse-facing Excel with
product photo, Coupang barcode, registered product name and work instructions.
"""
from __future__ import annotations

import io
import re
from difflib import SequenceMatcher
from typing import Any

PAGE_LABEL = "쿠팡 라벨 작업지시"


def _text(v: Any) -> str:
    return "" if v is None else str(v).strip()


def _oid(v: Any) -> str:
    s = _text(v)
    if s.upper().startswith("CP-"):
        s = s[3:]
    if s.endswith(".0") and s[:-2].isdigit():
        s = s[:-2]
    return s


def _num(v: Any) -> float:
    try:
        if v is None or _text(v) == "":
            return 0.0
        return float(str(v).replace(",", "").strip())
    except Exception:
        return 0.0


def _compact(v: Any) -> str:
    return re.sub(r"[^0-9A-Za-z가-힣]+", "", _text(v).lower())


def _tokens(v: Any) -> set[str]:
    s = re.sub(r"[^0-9A-Za-z가-힣]+", " ", _text(v).lower())
    stop = {
        "개", "세트", "상품", "제품", "구매", "쿠팡", "라벨", "부착", "포장",
        "블랙", "화이트", "실버", "그레이", "투명", "레드", "오렌지",
        "s", "m", "l", "xl", "free", "jd",
    }
    return {x for x in s.split() if len(x) >= 2 and x not in stop}


def _similarity(a: str, b: str) -> float:
    ca, cb = _compact(a), _compact(b)
    if not ca or not cb:
        return 0.0
    ta, tb = _tokens(a), _tokens(b)
    overlap = len(ta & tb) / max(1, len(ta | tb))
    contain = 1.0 if (len(ca) >= 4 and ca in cb) or (len(cb) >= 4 and cb in ca) else 0.0
    seq = SequenceMatcher(None, ca, cb).ratio()
    return min(1.0, contain * 0.42 + overlap * 0.38 + seq * 0.20)


def _header_key(v: Any) -> str:
    return re.sub(r"[\s_()\[\]·/\\-]+", "", _text(v).lower())


def _find_header(ws) -> tuple[int, dict[str, int]]:
    best = None
    for r in range(1, min(int(ws.max_row or 1), 30) + 1):
        values = [_header_key(ws.cell(r, c).value) for c in range(1, int(ws.max_column or 1) + 1)]
        cols: dict[str, int] = {}
        for c, key in enumerate(values, start=1):
            if not key:
                continue
            if "수량" in key and "금액" not in key:
                cols.setdefault("qty", c)
            if any(x in key for x in ("상품명", "제품명", "품명", "한글명", "상품이름")):
                cols.setdefault("name", c)
            if any(x in key for x in ("작업요청", "작업요청사항", "작업사항", "포장요청", "포장", "비고")):
                cols.setdefault("instruction", c)
            if any(x in key for x in ("주문번호", "발주번호", "오더번호")):
                cols.setdefault("order", c)
        score = (3 if "qty" in cols else 0) + (3 if "name" in cols else 0) + len(cols)
        if best is None or score > best[0]:
            best = (score, r, cols)
    if best is None or best[0] < 3:
        raise ValueError("발주서에서 수량/상품명 헤더를 찾지 못했습니다.")
    return best[1], best[2]


def _guess_name_col(ws, header_row: int, qty_col: int) -> int:
    best = None
    for c in range(1, int(ws.max_column or 1) + 1):
        if c == qty_col:
            continue
        hangul = 0
        nonempty = 0
        for r in range(header_row + 1, min(int(ws.max_row or 1), header_row + 60) + 1):
            s = _text(ws.cell(r, c).value)
            if not s:
                continue
            nonempty += 1
            if re.search(r"[가-힣]", s):
                hangul += 1
        score = hangul * 3 + nonempty
        if best is None or score > best[0]:
            best = (score, c)
    if not best or best[0] <= 0:
        raise ValueError("발주서에서 한글 상품명 열을 찾지 못했습니다.")
    return int(best[1])


def _image_map(ws) -> dict[int, bytes]:
    out: dict[int, bytes] = {}
    for img in list(getattr(ws, "_images", []) or []):
        try:
            row = int(img.anchor._from.row) + 1
            data = bytes(img._data())
            if data and row not in out:
                out[row] = data
        except Exception:
            continue
    return out


def parse_purchase_workbook(uploaded) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    raw = uploaded.getvalue() if hasattr(uploaded, "getvalue") else uploaded.read()
    if not raw:
        raise ValueError("업로드한 발주서가 비어 있습니다.")
    wb = load_workbook(io.BytesIO(raw), data_only=True)
    try:
        ws = wb.active
        header_row, cols = _find_header(ws)
        qty_col = cols.get("qty")
        if not qty_col:
            raise ValueError("발주서에서 수량 열을 찾지 못했습니다.")
        name_col = cols.get("name") or _guess_name_col(ws, header_row, qty_col)
        inst_col = cols.get("instruction")
        order_col = cols.get("order")
        images = _image_map(ws)

        rows = []
        for r in range(header_row + 1, int(ws.max_row or 1) + 1):
            name = _text(ws.cell(r, name_col).value)
            qty = _num(ws.cell(r, qty_col).value)
            if not name or qty <= 0:
                continue
            rows.append({
                "source_row": r,
                "order_no": _text(ws.cell(r, order_col).value) if order_col else "",
                "source_name": name,
                "purchase_qty": qty,
                "instruction": _text(ws.cell(r, inst_col).value) if inst_col else "",
                "image_bytes": images.get(r),
            })
        if not rows:
            raise ValueError("상품명과 수량이 함께 입력된 발주 행을 찾지 못했습니다.")
        return rows
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _load_erp(core) -> tuple[list[dict], dict[str, dict]]:
    barcode_mod = __import__("rg_barcode_print_v09193")
    barcode_mod.ensure_schema(core)
    with core._conn(core.DEFAULT_DB) as c:
        products = [dict(r) for r in c.execute(
            """SELECT id,item_code,option_id,name,item_type,active
               FROM products ORDER BY id DESC"""
        ).fetchall()]
        bom = [dict(r) for r in c.execute(
            """SELECT b.parent_product_id,b.component_product_id,b.qty_per,
                      p.name AS component_name,p.item_code AS component_code
               FROM bom_items b JOIN products p ON p.id=b.component_product_id"""
        ).fetchall()]
        masters = {
            _oid(r["vendor_item_id"]): dict(r)
            for r in c.execute(
                """SELECT vendor_item_id,barcode,product_name,option_name
                   FROM rg_barcode_master"""
            ).fetchall()
        }

    by_parent: dict[int, list[dict]] = {}
    for b in bom:
        by_parent.setdefault(int(b["parent_product_id"]), []).append(b)

    finished = []
    by_oid = {}
    for p in products:
        oid = _oid(p.get("option_id"))
        if not oid or str(p.get("item_type") or "") != "finished" or not int(p.get("active") or 0):
            continue
        comps = by_parent.get(int(p["id"]), [])
        saved = masters.get(oid, {})
        registered_name = _text(saved.get("product_name")) or _text(p.get("name"))
        option_name = _text(saved.get("option_name"))
        display_name = registered_name
        if option_name:
            reg_key = _compact(registered_name)
            opt_key = _compact(option_name)
            if opt_key and opt_key not in reg_key:
                display_name = f"{registered_name}, {option_name}" if registered_name else option_name
        # If the ERP finished-product name is already more complete than the
        # barcode-master split fields, keep the longer exact sellable name.
        erp_name = _text(p.get("name"))
        if len(_compact(erp_name)) > len(_compact(display_name)):
            display_name = erp_name

        row = {
            "id": int(p["id"]),
            "option_id": oid,
            "name": display_name,
            "registered_name": registered_name,
            "option_name": option_name,
            "display_name": display_name,
            "barcode": _text(saved.get("barcode")),
            "components": comps,
        }
        finished.append(row)
        by_oid[oid] = row
    return finished, by_oid


def _product_choices(core) -> tuple[list[str], dict[str, str]]:
    finished, _by_oid = _load_erp(core)
    choices = [""]
    choice_to_oid: dict[str, str] = {}
    for p in sorted(finished, key=lambda x: (x.get("name") or "", x.get("option_id") or "")):
        label = f"{p['option_id']} | {p['name']}"
        choices.append(label)
        choice_to_oid[label] = p["option_id"]
    return choices, choice_to_oid


def _rank_for_source(source_name: str, finished: list[dict]) -> list[tuple[float, dict, float]]:
    ranked = []
    for p in finished:
        best_component = 0.0
        best_qty = 1.0
        for c in p.get("components") or []:
            score = _similarity(source_name, _text(c.get("component_name")))
            if score > best_component:
                best_component = score
                best_qty = max(1.0, _num(c.get("qty_per")) or 1.0)
        product_score = _similarity(source_name, p.get("name") or "")
        score = max(best_component, product_score * 0.88)
        ranked.append((score, p, best_qty))
    ranked.sort(key=lambda x: (-x[0], x[1].get("name") or ""))
    return ranked[:5]


def prepare_rows(core, source_rows: list[dict]) -> list[dict]:
    finished, _by_oid = _load_erp(core)
    out = []
    for src in source_rows:
        ranked = _rank_for_source(src["source_name"], finished)
        top_score, top, qty_per = ranked[0] if ranked else (0.0, None, 1.0)
        second = ranked[1][0] if len(ranked) > 1 else 0.0
        auto = bool(top and top_score >= 0.52 and (top_score - second >= 0.06 or top_score >= 0.72))
        oid = top["option_id"] if auto else ""
        barcode = top.get("barcode", "") if auto else ""
        rg_name = top.get("name", "") if auto else ""
        label_qty = src["purchase_qty"] / qty_per if auto and qty_per > 0 else src["purchase_qty"]
        if abs(label_qty - round(label_qty)) < 1e-8:
            label_qty = int(round(label_qty))
        candidates = " | ".join(
            f"{p['option_id']} {p['name']} ({score:.2f})" for score, p, _q in ranked[:3]
        )
        out.append({
            "포함": True,
            "원본행": src["source_row"],
            "주문번호": src["order_no"],
            "발주상품명": src["source_name"],
            "발주수량": src["purchase_qty"],
            "쿠팡상품 선택": f"{oid} | {rg_name}" if oid and rg_name else "",
            "옵션ID": oid,
            "쿠팡등록상품명": rg_name,
            "바코드": barcode,
            "BOM구성수량": qty_per if auto else 1,
            "라벨수량": label_qty,
            "작업요청사항": src["instruction"],
            "추천후보": candidates,
            "상태": "자동매칭" if auto else "옵션ID 확인 필요",
            "image_bytes": src.get("image_bytes"),
        })
    return out


def _resolve_editor_rows(core, edited_rows: list[dict], prepared_rows: list[dict]) -> tuple[list[dict], list[str]]:
    _finished, by_oid = _load_erp(core)
    source_map = {int(r["원본행"]): r for r in prepared_rows}
    result, errors = [], []
    for row in edited_rows:
        if not bool(row.get("포함", True)):
            continue
        oid = _oid(row.get("옵션ID"))
        if not oid or oid not in by_oid:
            errors.append(f"{_text(row.get('발주상품명'))}: 위 '쿠팡 상품 매칭'에서 올바른 상품을 선택해 주세요.")
            continue
        product = by_oid[oid]
        barcode = _text(row.get("바코드")) or _text(product.get("barcode"))
        if not barcode:
            errors.append(f"{product.get('name')}: 쿠팡 바코드가 ERP에 없습니다.")
            continue
        label_qty = _num(row.get("라벨수량"))
        if label_qty <= 0 or abs(label_qty - round(label_qty)) > 1e-8:
            errors.append(f"{product.get('name')}: 라벨수량은 1 이상의 정수여야 합니다.")
            continue
        source_row = int(_num(row.get("원본행")) or 0)
        src = source_map.get(source_row, {})
        result.append({
            "order_no": _text(row.get("주문번호")),
            "source_name": _text(row.get("발주상품명")),
            "purchase_qty": _num(row.get("발주수량")),
            "option_id": oid,
            "product_name": _text(row.get("쿠팡등록상품명")) or _text(product.get("display_name")) or _text(product.get("name")),
            "coupang_full_option": _text(product.get("display_name")) or _text(row.get("쿠팡등록상품명")) or _text(product.get("name")),
            "barcode": barcode,
            "label_qty": int(round(label_qty)),
            "instruction": _text(row.get("작업요청사항")),
            "image_bytes": src.get("image_bytes"),
        })
    if not result and not errors:
        errors.append("작업지시서에 포함할 상품이 없습니다.")
    return result, errors


def _label_font(size: int, bold: bool = False):
    """Use Windows Malgun Gothic when available so Korean text stays sharp."""
    from pathlib import Path
    from PIL import ImageFont

    candidates = []
    if bold:
        candidates += [
            r"C:\Windows\Fonts\malgunbd.ttf",
            r"C:\Windows\Fonts\malgunbd.TTF",
        ]
    candidates += [
        r"C:\Windows\Fonts\malgun.ttf",
        r"C:\Windows\Fonts\malgun.TTF",
        r"C:\Windows\Fonts\gulim.ttc",
    ]
    for path in candidates:
        try:
            if Path(path).exists():
                return ImageFont.truetype(path, size=size)
        except Exception:
            pass
    try:
        return ImageFont.truetype("malgun.ttf", size=size)
    except Exception:
        return ImageFont.load_default()


def _label_50x30_png(barcode: str, product_name: str, scale: int = 4) -> io.BytesIO:
    """High-resolution preview of the ERP's real 50x30 direct-print label.

    Geometry follows rg_barcode_direct_print_v09257 exactly (400x240 base canvas)
    but is rendered at 4x resolution so Excel scaling does not blur barcode/text.
    """
    from PIL import Image, ImageDraw

    direct = __import__("rg_barcode_direct_print_v09257")
    barcode_mod = __import__("rg_barcode_print_v09193")

    base_w = int(direct.LABEL_W)
    base_h = int(direct.LABEL_H)
    W, H = base_w * scale, base_h * scale

    img = Image.new("RGB", (W, H), "white")
    draw = ImageDraw.Draw(img)
    code = _text(barcode)
    product = _text(product_name)

    # Same Code128 geometry used by direct printer.
    values = barcode_mod._code128_values(code)
    patterns = barcode_mod._CODE128_PATTERNS
    modules = 20 + sum(sum(int(ch) for ch in patterns[v]) for v in values)
    left, right = 16 * scale, (base_w - 16) * scale
    bar_top, bar_bottom = 8 * scale, 128 * scale
    usable = right - left
    module_scale = usable / modules
    x = left + 10 * module_scale

    for val in values:
        bar = True
        for ch in patterns[val]:
            mw = int(ch) * module_scale
            if bar:
                x0 = int(round(x))
                x1 = max(x0 + 1, int(round(x + mw)))
                draw.rectangle([x0, bar_top, x1 - 1, bar_bottom - 1], fill="black")
            x += mw
            bar = not bar

    def _center_text(text: str, y_base: int, px_base: int):
        if not text:
            return
        font = _label_font(px_base * scale)
        box = draw.textbbox((0, 0), text, font=font)
        tw = box[2] - box[0]
        draw.text(((W - tw) / 2, y_base * scale), text, font=font, fill="black")

    # Barcode number: same y/size as ERP direct print.
    _center_text(code, 134, 18)

    # Product name: same two-line wrapping principle as direct print.
    product_font = _label_font(24 * scale)
    maxw = (base_w - 24) * scale

    def _measure(text: str) -> int:
        box = draw.textbbox((0, 0), text, font=product_font)
        return box[2] - box[0]

    def _wrap_two_lines(text: str) -> list[str]:
        text = _text(text)
        if not text:
            return [""]
        words = text.split() or [text]
        lines = []
        current = ""
        for word in words:
            trial = word if not current else current + " " + word
            if _measure(trial) <= maxw:
                current = trial
            else:
                if current:
                    lines.append(current)
                    current = word
                else:
                    chunk = ""
                    for ch in word:
                        if _measure(chunk + ch) <= maxw:
                            chunk += ch
                        else:
                            if chunk:
                                lines.append(chunk)
                            chunk = ch
                    current = chunk
            if len(lines) >= 2:
                break
        if len(lines) < 2 and current:
            lines.append(current)
        return lines[:2] or [text]

    lines = _wrap_two_lines(product)
    if len(lines) <= 1:
        _center_text(lines[0] if lines else "", 164, 24)
    else:
        _center_text(lines[0], 150, 24)
        _center_text(lines[1], 178, 24)

    _center_text("MADE IN CHINA", 210, 24)

    out = io.BytesIO()
    img.save(out, format="PNG", dpi=(300, 300), optimize=False)
    out.seek(0)
    return out


def _sample_label_png() -> io.BytesIO:
    """Fixed crossed-out sample label shown in the warehouse instruction template."""
    from PIL import Image, ImageDraw

    base = _label_50x30_png("S0000KCT828282", "상품 옵션명 / 옵션 1개", scale=4)
    img = Image.open(base).convert("RGB")
    draw = ImageDraw.Draw(img)
    w, h = img.size
    red = (230, 90, 90)
    draw.line((w * 0.18, h * 0.20, w * 0.82, h * 0.72), fill=red, width=max(4, w // 180))
    draw.line((w * 0.82, h * 0.20, w * 0.18, h * 0.72), fill=red, width=max(4, w // 180))
    out = io.BytesIO()
    img.save(out, format="PNG")
    out.seek(0)
    return out


def build_instruction_xlsx(rows: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "쿠팡 라벨 작업지시"

    # Exact column structure of the user's existing warehouse worksheet:
    # 주문번호 / 옵션 / 사진 / 수량 / 바코드번호 / 라벨이미지 / 작업요청사항
    headers = [
        "주문번호",
        "옵션",
        "사진",
        "수량",
        "바코드",
        "바코드 및 한글표시스티커\n"
        "바코드는 사진을 크게 캡처해서 아래 사진보다\n"
        "선명한 사진으로 넣어주세요. 흐린사진은 바코드 인식을\n"
        "못해서 추가 PDF 파일로 전달해주셔도 됩니다.",
        "작업요청사항",
    ]
    for col, value in enumerate(headers, start=1):
        ws.cell(1, col, value)

    header_fill = PatternFill("solid", fgColor="DDD8C0")
    sample_fill = PatternFill("solid", fgColor="FFF200")
    thin = Side(style="thin", color="555555")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col in range(1, 8):
        cell = ws.cell(1, col)
        cell.font = Font(name="맑은 고딕", bold=True, size=10)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[1].height = 62

    # Match the reference proportions closely.
    for col, width in {
        "A": 12.5,
        "B": 39,
        "C": 18,
        "D": 12.5,
        "E": 18,
        "F": 34,
        "G": 52,
    }.items():
        ws.column_dimensions[col].width = width

    # Fixed sample row, kept as part of the template.
    sample_row = 2
    ws.cell(sample_row, 1, "241004012\n\n샘플")
    ws.cell(sample_row, 2, "상품 옵션명 / 옵션 1개")
    ws.cell(sample_row, 3, "제품사진")
    ws.cell(sample_row, 4, 320)
    ws.cell(sample_row, 5, "S0000KCT828282")
    ws.cell(sample_row, 6, "")
    ws.cell(
        sample_row,
        7,
        "개별 박스 위에 바코드 스티커 부착부탁드립니다.\n"
        "20Kg 미만 160cm 미만 포장으로 해주세요.",
    )
    for col in range(1, 8):
        cell = ws.cell(sample_row, col)
        cell.font = Font(name="맑은 고딕", size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.cell(sample_row, 1).fill = sample_fill
    ws.cell(sample_row, 1).font = Font(name="맑은 고딕", bold=True, size=10)
    ws.row_dimensions[sample_row].height = 105

    try:
        simg = XLImage(_sample_label_png())
        simg.width = 195
        simg.height = 117
        ws.add_image(simg, f"F{sample_row}")
    except Exception:
        pass

    data_start = 3
    for idx, row in enumerate(rows, start=data_start):
        ws.cell(idx, 1, row.get("order_no") or "")
        ws.cell(idx, 2, row.get("coupang_full_option") or row.get("product_name") or "")
        ws.cell(idx, 4, int(row.get("label_qty") or 0))
        ws.cell(idx, 5, _text(row.get("barcode")))
        ws.cell(idx, 6, "")
        ws.cell(idx, 7, row.get("instruction") or "")

        for col in range(1, 8):
            cell = ws.cell(idx, col)
            cell.font = Font(name="맑은 고딕", size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border

        # Reference sheet uses tall product rows.
        ws.row_dimensions[idx].height = 112

        img_bytes = row.get("image_bytes")
        if img_bytes:
            try:
                photo = XLImage(io.BytesIO(img_bytes))
                photo.width = 112
                photo.height = 86
                ws.add_image(photo, f"C{idx}")
            except Exception:
                pass

        try:
            label_png = _label_50x30_png(
                _text(row.get("barcode")),
                _text(row.get("product_name")),
                scale=4,
            )
            bimg = XLImage(label_png)
            # Same visual size as the user's reference, but source remains 1600x960.
            bimg.width = 195
            bimg.height = 117
            ws.add_image(bimg, f"F{idx}")
        except Exception as exc:
            ws.cell(idx, 6, f"5x3 라벨 이미지 생성 실패: {exc}")

    # Merge repeated order numbers vertically, like the reference template.
    if rows:
        merge_start = data_start
        last_order = _text(ws.cell(data_start, 1).value)
        last_data_row = data_start + len(rows) - 1
        for r in range(data_start + 1, last_data_row + 2):
            order = _text(ws.cell(r, 1).value) if r <= last_data_row else "__END__"
            if order != last_order:
                if last_order and r - merge_start > 1:
                    ws.merge_cells(
                        start_row=merge_start, start_column=1,
                        end_row=r - 1, end_column=1,
                    )
                    ws.cell(merge_start, 1).alignment = Alignment(
                        horizontal="center", vertical="center", wrap_text=True
                    )
                merge_start = r
                last_order = order

    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.15
    ws.page_margins.right = 0.15
    ws.page_margins.top = 0.25
    ws.page_margins.bottom = 0.25
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _records(value) -> list[dict]:
    if hasattr(value, "to_dict"):
        return [dict(x) for x in value.to_dict("records")]
    return [dict(x) for x in (value or [])]


def render_page(st, core):
    import pandas as pd

    st.title(PAGE_LABEL)
    st.caption(
        "중국 배대지에 보낼 편집된 구매대행 발주 Excel을 업로드하면 ERP의 RG상품/BOM 및 쿠팡 바코드를 연결해 "
        "라벨 작업지시 Excel을 만듭니다. 자동생성된 매입자료를 쓰지 않고, 업로드한 발주서를 기준으로만 작업합니다."
    )
    st.info(
        "업로드 후 자동 매칭 결과를 먼저 확인할 수 있습니다. 옵션ID·라벨수량·작업요청사항은 출력 전에 직접 수정할 수 있습니다."
    )

    uploaded = st.file_uploader(
        "편집 완료한 중국 구매대행 발주서 Excel",
        type=["xlsx"],
        key="coupang_label_instruction_v09295_upload",
    )
    if uploaded is None:
        return

    try:
        source_rows = parse_purchase_workbook(uploaded)
        prepared = prepare_rows(core, source_rows)
    except Exception as exc:
        st.error(f"발주서 확인 실패: {exc}")
        return

    auto_count = sum(1 for r in prepared if r["상태"] == "자동매칭")
    c1, c2, c3 = st.columns(3)
    c1.metric("발주 품목", f"{len(prepared):,}개")
    c2.metric("자동 매칭", f"{auto_count:,}개")
    c3.metric("확인 필요", f"{len(prepared)-auto_count:,}개")

    choices, choice_to_oid = _product_choices(core)
    _finished, by_oid = _load_erp(core)

    st.subheader("쿠팡 상품 매칭")
    st.caption(
        "각 발주상품 오른쪽 선택상자를 눌러 연결할 쿠팡 상품을 직접 바꿀 수 있습니다. "
        "상품명 일부를 입력하면 목록을 빠르게 검색할 수 있습니다."
    )

    selected_rows = []
    for item in prepared:
        source_row = int(item["원본행"])
        default_choice = _text(item.get("쿠팡상품 선택"))
        try:
            default_index = choices.index(default_choice) if default_choice in choices else 0
        except Exception:
            default_index = 0

        col_name, col_select = st.columns([1.05, 2.2], vertical_alignment="center")
        col_name.markdown(
            f"**{item['발주상품명']}**  \\n"
            f"발주수량 {item['발주수량']:g} · 원본행 {source_row}"
        )
        selected_choice = col_select.selectbox(
            f"{item['발주상품명']} 쿠팡상품",
            options=choices,
            index=default_index,
            key=f"coupang_label_instruction_v09295_match_{source_row}",
            label_visibility="collapsed",
            placeholder="쿠팡 상품을 선택하세요",
        )

        row = dict(item)
        row["쿠팡상품 선택"] = selected_choice or ""
        oid = choice_to_oid.get(selected_choice or "", "")
        if oid and oid in by_oid:
            product = by_oid[oid]
            row["옵션ID"] = oid
            row["쿠팡등록상품명"] = _text(product.get("name"))
            row["바코드"] = _text(product.get("barcode"))
            # Recalculate BOM 구성/라벨수량 from the selected product and source item.
            best_qty = 1.0
            best_score = 0.0
            for comp in product.get("components") or []:
                score = _similarity(item["발주상품명"], _text(comp.get("component_name")))
                if score > best_score:
                    best_score = score
                    best_qty = max(1.0, _num(comp.get("qty_per")) or 1.0)
            row["BOM구성수량"] = best_qty
            label_qty = _num(item["발주수량"]) / best_qty if best_qty > 0 else _num(item["발주수량"])
            if abs(label_qty - round(label_qty)) < 1e-8:
                label_qty = int(round(label_qty))
            row["라벨수량"] = label_qty
            row["상태"] = "선택 완료"
        else:
            row["옵션ID"] = ""
            row["쿠팡등록상품명"] = ""
            row["바코드"] = ""
            row["상태"] = "옵션ID 확인 필요"
        selected_rows.append(row)

    st.divider()
    st.subheader("출력 내용 확인 및 수정")
    st.caption("상품 매칭 결과가 아래에 즉시 반영됩니다. 라벨수량과 작업요청사항은 여기서 수정할 수 있습니다.")

    frame = pd.DataFrame([{k: v for k, v in r.items() if k != "image_bytes"} for r in selected_rows])
    columns = [
        "포함", "원본행", "주문번호", "발주상품명", "발주수량", "옵션ID",
        "쿠팡등록상품명", "바코드", "BOM구성수량", "라벨수량", "작업요청사항", "상태",
    ]
    frame = frame[[col for col in columns if col in frame.columns]]
    edited = st.data_editor(
        frame,
        use_container_width=True,
        hide_index=True,
        num_rows="fixed",
        height=min(760, max(300, 38 * (min(len(frame), 17) + 1))),
        key="coupang_label_instruction_v09295_editor",
        disabled=["원본행", "발주상품명", "발주수량", "옵션ID", "쿠팡등록상품명", "바코드", "상태", "BOM구성수량"],
        column_config={
            "포함": st.column_config.CheckboxColumn("포함", width="small"),
            "원본행": st.column_config.NumberColumn("원본행", width="small"),
            "발주상품명": st.column_config.TextColumn("발주상품명", width="large"),
            "발주수량": st.column_config.NumberColumn("발주수량", width="small"),
            "옵션ID": st.column_config.TextColumn("쿠팡 옵션ID", width="medium"),
            "쿠팡등록상품명": st.column_config.TextColumn("쿠팡 등록상품명", width="large"),
            "바코드": st.column_config.TextColumn("바코드", width="medium"),
            "BOM구성수량": st.column_config.NumberColumn("구성수량", width="small"),
            "라벨수량": st.column_config.NumberColumn("라벨수량", min_value=1, step=1, width="small"),
            "작업요청사항": st.column_config.TextColumn("작업요청사항", width="large"),
            "상태": st.column_config.TextColumn("상태", width="small"),
        },
    )

    rows, errors = _resolve_editor_rows(core, _records(edited), prepared)

    st.divider()
    st.subheader("작업지시서 출력")
    if errors:
        st.warning("아래 확인 필요 항목을 먼저 수정하면 출력 버튼이 활성화됩니다.")
        for msg in errors[:12]:
            st.caption("• " + msg)
        st.button(
            "쿠팡 라벨 작업지시서 만들기",
            disabled=True,
            use_container_width=True,
            key="coupang_label_instruction_v09295_generate_disabled",
        )
        return

    st.success(f"출력 준비 완료: {len(rows):,}개 품목 / 라벨 {sum(r['label_qty'] for r in rows):,}장")
    if st.button(
        "쿠팡 라벨 작업지시서 만들기",
        type="primary",
        use_container_width=True,
        key="coupang_label_instruction_v09295_generate",
    ):
        try:
            st.session_state["coupang_label_instruction_v09295_file"] = build_instruction_xlsx(rows)
        except Exception as exc:
            st.error(f"작업지시서 생성 실패: {exc}")
            return

    xlsx = st.session_state.get("coupang_label_instruction_v09295_file")
    if not xlsx:
        st.caption("위 버튼을 누르면 Excel 파일을 만든 뒤 다운로드 버튼이 나타납니다.")
        return

    base = re.sub(r"\.[Xx][Ll][Ss][Xx]$", "", getattr(uploaded, "name", "발주서"))
    st.download_button(
        "Excel 다운로드",
        data=xlsx,
        file_name=f"{base}_쿠팡라벨작업지시.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
        use_container_width=True,
        key="coupang_label_instruction_v09295_download",
    )
